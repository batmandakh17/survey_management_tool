"""Write the KoBoToolbox call-record form (XLSForm): formats/kobo/rtm_call_record.xlsx.

Run from the tool folder:  python3 formats/kobo/make_xlsform.py
One submission per call attempt. The household file households.csv, written by the leader view (Call log tab),
is the form's media file: it fills the household's name, numbers, planned time, program and round. Needs openpyxl.
XLSForm references: question types, select_one_from_file, choice_filter, pulldata(), dynamic defaults, start/end
metadata (KoBoToolbox support site and ODK XLSForm documentation).
"""
import os

from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "rtm_call_record.xlsx")

RESULTS = [("COMPLETE", "Interview completed"), ("APPOINTMENT", "Contact made, callback booked"), ("PARTIAL", "Interview started, not finished"),
           ("NOT_AVAILABLE", "Household reached, respondent not available"), ("NO_ANSWER", "No answer"), ("BUSY", "Busy"),
           ("SWITCHED_OFF", "Switched off or unreachable"), ("DEAD_NUMBER", "Number not in service"), ("WRONG_NUMBER", "Wrong number"),
           ("REFUSAL_SOFT", "Refusal, may retry"), ("REFUSAL_HARD", "Refusal, do not retry"), ("LANGUAGE", "Language barrier"),
           ("OTHER", "Other (write a note)"), ("NONE", "No call made: an action or a note only")]
CONTACT = ("COMPLETE", "APPOINTMENT", "PARTIAL", "NOT_AVAILABLE", "REFUSAL_SOFT", "REFUSAL_HARD", "LANGUAGE", "OTHER")   # the household was reached
WILLING = [("yes", "Yes, willing to take part"), ("later", "Later: asked to be called another time"), ("no", "No, does not want to take part"),
           ("not_asked", "Not asked")]
ACTIONS = [("none", "Keep calling"), ("RELEASE", "Release for this round"), ("SWITCH_REQUEST", "Ask the leader for another enumerator next round"),
           ("REMOVE_REQUEST", "Ask the leader to remove (no number works)"), ("ACTIVATE_REPLACEMENT", "Activate a replacement household")]
SLOTS = [("1", "Number 1: ${phone_1} (${phone_1_label})"), ("2", "Number 2: ${phone_2} (${phone_2_label})"), ("3", "Number 3: ${phone_3} (${phone_3_label})")]

SURVEY_COLS = ["type", "name", "label", "hint", "required", "appearance", "choice_filter", "parameters", "calculation", "default", "relevant", "constraint", "constraint_message"]


def pull(col):
    return f"pulldata('households', '{col}', 'hhid', ${{hhid}})"


def survey_rows():
    r = []
    add = lambda **k: r.append(k)
    add(type="start", name="start")
    add(type="end", name="end")
    add(type="deviceid", name="deviceid")
    add(type="text", name="enumerator_id", label="Your enumerator ID", hint="As written in your assignment (for example E07).", required="yes")
    add(type="select_one_from_file households.csv", name="hhid", label="Household", hint="Your households for this round, and their replacements once activated. Type part of the ID to search.",
        required="yes", appearance="autocomplete", choice_filter="enumerator_id = ${enumerator_id}", parameters="value=hhid label=label")
    for col in ("program", "round", "contact_name", "phone_1", "phone_1_label", "phone_2", "phone_2_label", "phone_3", "phone_3_label",
                "planned_date", "planned_hour", "replaces", "psu", "stratum"):
        add(type="calculate", name=col, calculation=pull(col))
    add(type="note", name="who", label="${contact_name} · planned ${planned_date} at ${planned_hour}:00 · PSU ${psu}",
        hint="The numbers are listed below. Dial outside the tool and record the call here.")
    add(type="select_one slots", name="phone_slot", label="Which number did you dial?", required="${result} != 'NONE'",
        choice_filter="(name = '1' and ${phone_1} != '') or (name = '2' and ${phone_2} != '') or (name = '3' and ${phone_3} != '')")
    add(type="datetime", name="call_start", label="Call started", required="${result} != 'NONE'", default="now()")
    add(type="datetime", name="call_end", label="Call ended", relevant="${result} != 'NONE'",
        constraint=". >= ${call_start}", constraint_message="The call cannot end before it started.")
    add(type="select_one results", name="result", label="Result of the call", required="yes")
    add(type="select_one willing", name="willing", label="Willing to take part?", required="yes",
        relevant=" or ".join(f"${{result}} = '{c}'" for c in CONTACT))
    add(type="datetime", name="appointment", label="Callback agreed for", relevant="${result} = 'APPOINTMENT' or ${willing} = 'later'",
        constraint=". >= ${call_start}", constraint_message="The callback must be after this call.")
    add(type="select_one actions", name="action", label="After this call", required="yes", default="none")
    add(type="select_one_from_file households.csv", name="target_hhid", label="Replacement to activate", hint="The household's replacements, in the order they are to be used.",
        relevant="${action} = 'ACTIVATE_REPLACEMENT'", required="${action} = 'ACTIVATE_REPLACEMENT'", appearance="minimal",
        choice_filter="kind = 'replacement' and replaces = ${hhid}", parameters="value=hhid label=label")
    add(type="text", name="note", label="Note", hint="Operational only: no personal details.")
    return r


def main():
    wb = Workbook()
    s = wb.active
    s.title = "survey"
    s.append(SURVEY_COLS)
    for row in survey_rows():
        s.append([row.get(c, "") for c in SURVEY_COLS])
    c = wb.create_sheet("choices")
    c.append(["list_name", "name", "label"])
    for ln, items in (("results", RESULTS), ("willing", WILLING), ("actions", ACTIONS), ("slots", SLOTS)):
        for name, label in items:
            c.append([ln, name, label])
    st = wb.create_sheet("settings")
    st.append(["form_title", "form_id", "version", "default_language", "instance_name"])
    st.append(["RTM call record", "rtm_call_record", "2026110101", "English (en)", "concat(${enumerator_id}, ' ', ${hhid}, ' ', ${result})"])
    wb.save(OUT)
    print(f"written {OUT}")


if __name__ == "__main__":
    main()
