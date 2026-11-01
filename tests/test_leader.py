"""End-to-end test of the leader view (Step 3), in headless Chrome, on the synthetic samples.

Run from the tool folder:  python3 tests/test_leader.py
Writes tests/results_leader.txt and screenshots in tests/screens/leader_*.png.
The analytics are checked against an independent Python implementation of the handbook's rules (tests/ref_stats.py).
"""
import csv
import glob
import hashlib
import io
import json
import os
import sys
import time
from contextlib import redirect_stdout

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(ROOT, "samples"))
sys.path.insert(0, os.path.join(ROOT, "formats"))
from cdp import Browser  # noqa: E402
from make_protected_sample import unseal as py_unseal  # noqa: E402
import ref_stats  # noqa: E402

SCRATCH = os.environ.get("RTM_SCRATCH", os.path.join(HERE, "_work"))
SCREENS = os.path.join(HERE, "screens")
os.makedirs(SCREENS, exist_ok=True)
STAMP = time.strftime("%H%M%S")
DL = os.path.join(SCRATCH, f"ldl_{STAMP}")
TOOL = "file://" + os.path.join(ROOT, "rtm_interview_tool.html")
S = os.path.join(ROOT, "samples", "leader")
HIST = sorted(glob.glob(os.path.join(S, "history", "*")))
EXPORTS = sorted(glob.glob(os.path.join(S, "exports_R05", "*")))
MANIFEST = os.path.join(S, "L2MN-DEMO_R05_manifest.json")
results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok), detail))
    print(("PASS " if ok else "FAIL ") + name + (f"  [{str(detail)[:600]}]" if detail and not ok else ""))


def canon(o):
    return json.dumps(o, sort_keys=True, separators=(",", ":"), ensure_ascii=False)


def sha(s):
    return hashlib.sha256(s.encode("utf-8")).hexdigest()


def snapshot():
    return {p: os.path.getmtime(p) for p in glob.glob(os.path.join(DL, "*"))}


def wait_new(pattern, before, timeout=20):
    t0 = time.time()
    while time.time() - t0 < timeout:
        for p in sorted(glob.glob(os.path.join(DL, pattern))):
            if not p.endswith(".crdownload") and (p not in before or os.path.getmtime(p) > before[p]):
                time.sleep(0.3)
                return p
        time.sleep(0.2)
    return None


def main():
    b = Browser(os.path.join(SCRATCH, f"lprofile_{STAMP}"), port=9344)
    try:
        b.ws.call("Emulation.setTimezoneOverride", timezoneId="Asia/Ulaanbaatar")
        b.downloads_to(DL)
        b.goto(TOOL)
        b.js("RTM.setNow('2026-10-23T15:00:00+08:00'); APP.noFSA = true; true")

        def setv(sel, v):
            b.js(f"(() => {{ const el = document.querySelector({json.dumps(sel)}); el.value = {json.dumps(v)}; el.dispatchEvent(new Event('change', {{bubbles: true}})); return true; }})()")

        def tick(sel, on):
            b.js(f"(() => {{ const el = document.querySelector({json.dumps(sel)}); el.checked = {str(bool(on)).lower()}; el.dispatchEvent(new Event('change', {{bubbles: true}})); return true; }})()")

        def start(leader="L01", program="L2MN-DEMO", country="MN", rnd=5):
            if b.js("!!(window.LEAD && LEAD.state())"):
                b.click('[data-lact="close"]'); b.wait("APP.modal"); b.js("(document.querySelector('[data-act=\"modalalt\"]') || document.querySelector('[data-act=\"modalok\"]')).click(); true"); b.wait("!LEAD.state()")
            if b.js("APP.view") != "leader":
                b.click('[data-act="view"][data-v="leader"]')
            b.wait("!!document.querySelector('#lstart')")
            b.js(f"""(() => {{ const f = document.querySelector('#lstart'); f.querySelector('[name=leader]').value = {json.dumps(leader)};
                   f.querySelector('[name=program]').value = {json.dumps(program)}; f.querySelector('[name=country]').value = {json.dumps(country)};
                   f.querySelector('[name=round]').value = {json.dumps(str(rnd))}; return true; }})()""")
            b.click('[data-lact="start"]'); b.wait("!!LEAD.state()")

        # ================= A. preparing round 5 =================
        b.click('[data-act="view"][data-v="leader"]')
        b.wait("!!document.querySelector('#lstart')")
        b.shot(os.path.join(SCREENS, "leader_01_start.png"))
        check("the leader view states plainly that it has no server, shared database or network",
              "no server, no shared database and no network connection" in b.js("document.body.innerText"))
        start()
        b.click('[data-lact="apply"]')
        errs = b.js("LEAD.ui.errors")
        check("applying the rule with nothing loaded lists what is missing", any("household list" in e for e in errs) and any("rotation share" in e for e in errs), errs)
        b.set_files("#lf-hh", [os.path.join(S, "household_list.csv")]); b.wait("!!LEAD.state().inputs.households")
        b.set_files("#lf-hist", HIST); b.wait(f"LEAD.state().inputs.history.length === {len(HIST)}")
        b.set_files("#lf-repl", [os.path.join(S, "selection.csv")]); b.wait("!!LEAD.state().inputs.replacements")
        txt = b.js("document.querySelector('.panel').innerText")
        check("the household list is read: 110 to assign, 34 unused replacements, 6 removed and kept", "110 households to assign, 34 unused replacements, 6 removed" in txt, txt[:600])
        # interviewed per round, against the Python generator's count (Box 3.0 rule)
        py_int = ref_stats.interviewed_per_round_strict(HIST)
        js_int = b.js("[...LEAD.state() && LCORE.buildHistory(LEAD.state().inputs.history.map(f => LCORE.readHistoryFile(f.name, f.text)), {tz: 'Asia/Ulaanbaatar'}).complete].map(([r, m]) => [r, m.size])")
        check("earlier rounds: interviewed households per round equal an independent count from the calls files", [list(x) for x in js_int] == py_int, [js_int, py_int])
        check("the rtmtools selection is read as is: 36 replacements dealt to 36 households, round-robin per PSU", "rtmtools selection: 36 replacements, one for each of 36 households" in txt)
        check("the team, dial cap and calendar habits are taken from the last manifest", b.js("LEAD.state().team.map(t => t.id).join(',')") == "E01,E02,E03,E07" and b.js("LEAD.state().cal.dial_cap") == 8)
        setv('[data-lset="cal.window_start"]', "2026-10-26"); setv('[data-lset="cal.window_end"]', "2026-11-20")
        check("the calendar lists the 10 November holiday from the bundled table", "Genghis Khan" in b.js("document.querySelector('.panel').innerText"))
        tick('[data-lteam="2"][data-f="available"]', False)            # E03 away
        b.click('[data-lact="apply"]')
        check("keep the same enumerator refuses to run without this round's rotation share (no default, D4)", any("no default" in e for e in b.js("LEAD.ui.errors")))
        setv('[data-lset="rule.rotation_share"]', "10"); setv('[data-lset="rule.seed"]', "20261027")
        b.click('[data-lact="apply"]'); b.wait("!!LEAD.state().plan")
        b.shot(os.path.join(SCREENS, "leader_02_prepare.png"), full=True)
        P = json.loads(b.js("JSON.stringify([...LEAD.state().plan.out])"))
        out = {h: x for h, x in P}
        sums = json.loads(b.js("JSON.stringify([...LEAD.state().plan.sums].map(([h, s]) => [h, s]))"))
        sums = {h: s for h, s in sums}
        check("every household to assign has an enumerator, none of them the one away", len(out) == 110 and all(x["enumerator_id"] != "E03" for x in out.values()))
        away = [h for h, x in out.items() if x["reason"] == "enumerator_unavailable"]
        check("the absent enumerator's households are dealt to the others, reason 'enumerator unavailable'", away and all(out[h].get("previous_enumerator") == "E03" for h in away), len(away))
        keep_ok = all(x["enumerator_id"] == ((sums[h].get("last_interview") or {}).get("enumerator") or sums[h].get("previous")) for h, x in out.items() if x["reason"] == "keep")
        check("kept households stay with the enumerator who completed their last interview (step 1)", keep_ok)
        hard = [h for h, x in out.items() if x["reason"] == "switch_hard_to_reach"]
        hard_def = all((not sums[h]["interviewed_last"] and sums[h]["in_last"] and sums[h]["dials_last"] >= 8) or sums[h]["switch_requested"] for h in hard)
        check("hard-to-reach households (full cap missed, or a switch request) go to another enumerator (D5)",
              hard and hard_def and all(out[h]["enumerator_id"] != out[h]["previous_enumerator"] for h in hard), hard)
        rot = [h for h, x in out.items() if x["reason"] == "rotation"]
        kept_before = {}
        for h, x in out.items():
            if x["reason"] in ("keep", "rotation"):
                e0 = x.get("previous_enumerator") if x["reason"] == "rotation" else x["enumerator_id"]
                kept_before[e0] = kept_before.get(e0, 0) + 1
        given = {}
        for h in rot:
            given[out[h]["previous_enumerator"]] = given.get(out[h]["previous_enumerator"], 0) + 1
        check("rotation moves round(10%) of each enumerator's kept households to another enumerator (step 4)",
              all(given.get(e, 0) == round(0.10 * n) for e, n in kept_before.items()) and all(out[h]["enumerator_id"] != out[h]["previous_enumerator"] for h in rot), [given, kept_before])
        new = [h for h, x in out.items() if x["reason"] == "new"]
        check("households new to the panel are dealt at random, evenly (2 each)", sorted([out[h]["enumerator_id"] for h in new]) == ["E01", "E01", "E02", "E02", "E07", "E07"], new)
        # reproducibility
        b.click('[data-lact="apply"]'); time.sleep(0.3)
        check("the same inputs and seed give the same assignment", json.loads(b.js("JSON.stringify([...LEAD.state().plan.out])")) == P)
        setv('[data-lset="rule.seed"]', "777"); b.click('[data-lact="apply"]'); time.sleep(0.3)
        check("another seed gives another rotation", json.loads(b.js("JSON.stringify([...LEAD.state().plan.out])")) != P)
        # review cycle 5: a setting changed after the rule was applied sets the assignment aside until it is applied again
        seed_now = b.js("LEAD.state().rule.seed")
        setv('[data-lset="rule.seed"]', "778")
        aside = b.js("[LEAD.state().plan === null, document.querySelector('.panel').innerText.includes('Apply the rule again')]")
        setv('[data-lset="rule.seed"]', str(seed_now)); b.click('[data-lact="apply"]'); b.wait("!!LEAD.state().plan")
        check("a setting changed after the rule was applied (here the seed) sets the assignment aside and asks for the rule again",
              aside == [True, True] and not b.js("document.querySelector('.panel').innerText.includes('Apply the rule again')"), aside)
        setv('[data-lset="rule.seed"]', "20261027"); b.click('[data-lact="apply"]'); time.sleep(0.3)
        # load check and balance
        loads = b.js("LCORE.loads(Object.assign({}, {team: LEAD.state().team, cap: 8}), LEAD.state().plan.out, LEAD.state().plan.sums)")
        check("the load check shows each enumerator's expected dials against the target", len(loads) == 3 and any(abs(l["departure"]) > 0.15 for l in loads), loads)
        b.click('[data-lact="balance"]'); time.sleep(0.3)
        loads2 = b.js("LCORE.loads(Object.assign({}, {team: LEAD.state().team, cap: 8}), LEAD.state().plan.out, LEAD.state().plan.sums)")
        moved = [x for x in b.js("[...LEAD.state().plan.out.values()]") if x["reason"] == "capacity"]
        check("balancing moves kept households with reason 'capacity' until every load is within the margin",
              moved and all(abs(l["departure"]) <= 0.15 + 1e-9 for l in loads2) and all(m.get("previous_enumerator") for m in moved), [loads2, len(moved)])
        # planning
        plans = json.loads(b.js("JSON.stringify([...LEAD.state().plan.plans])"))
        days = b.js("LCORE.workingDays(LEAD.state() && (() => { const c = Object.assign({}, LEAD.state().cal); c.holidays = [{date: '2026-11-10', working: false}]; return c; })())")
        first_part = days[: round(0.5 * len(days))]
        check("first dials fall on the first half of the working days, never a Sunday or the 10 November holiday",
              all(p["first_dial_date"] in first_part for _, p in plans) and "2026-11-10" not in [p["first_dial_date"] for _, p in plans], first_part)
        same_hour = [h for h, p in plans if sums[h].get("last_interview") and sums[h]["last_interview"].get("hour") is not None]
        check("each household's first dial is planned at the hour of its last interview (3.5.2)",
              same_hour and all(dict(plans)[h]["target_hour"] == sums[h]["last_interview"]["hour"] for h in same_hour))
        # write the files
        before = snapshot()
        b.click('[data-lact="write"]'); b.wait("!!LEAD.ui.ready", timeout=30)
        b.click('[data-lact="dlall"]')
        got = [wait_new(n, before) for n in ("E01_assignment_R05.json", "E02_assignment_R05.json", "E07_assignment_R05.json", "L2MN-DEMO_R05_manifest.json", "calendar.csv")]
        check("one assignment file per available enumerator, the manifest and the calendar row are written", all(got), got)
        A = {e: json.load(open(os.path.join(DL, f"{e}_assignment_R05.json"), encoding="utf-8")) for e in ("E01", "E02", "E07")}
        M = json.load(open(os.path.join(DL, "L2MN-DEMO_R05_manifest.json"), encoding="utf-8"))
        hl = list(csv.DictReader(open(os.path.join(S, "household_list.csv"), encoding="utf-8")))
        phones_of = {r["hhid"]: {r[k] for k in ("phone_1", "phone_2", "phone_3") if r.get(k)} for r in hl}
        leak = []
        for e, a in A.items():
            mine = {h["hhid"] for h in a["households"]} | {r["hhid"] for h in a["households"] for r in h.get("replacements", [])}
            allowed = set().union(*[phones_of[h] for h in mine])
            for h in a["households"]:
                for p in h["contact"]["phones"]:
                    if p["number"] not in allowed: leak.append((e, h["hhid"]))
                for r in h.get("replacements", []):
                    for p in r["contact"]["phones"]:
                        if p["number"] not in allowed: leak.append((e, r["hhid"]))
            others = set().union(*[phones_of[h] for h in phones_of if h not in mine])
            text = json.dumps(a, ensure_ascii=False)
            leak += [(e, n) for n in others if n and n in text]
        check("access control: no assignment file holds another enumerator's contact information", not leak, leak[:5])
        ids = [h["hhid"] for a in A.values() for h in a["households"]]
        reps = [r["hhid"] for a in A.values() for h in a["households"] for r in h.get("replacements", [])]
        check("each household is in exactly one file, and each replacement belongs to exactly one household", len(ids) == len(set(ids)) == 110 and len(reps) == len(set(reps)) == 34, [len(ids), len(reps)])
        check("the manifest holds IDs only (no names or phone numbers) and every file's fingerprint",
              "+000" not in json.dumps(M) and "Contact" not in json.dumps(M) and {e["enumerator_id"]: e["fingerprint"] for e in M["enumerators"]} == {e: a["fingerprint"] for e, a in A.items()})
        import check_formats as cf
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = cf.main([os.path.join(DL, f"{e}_assignment_R05.json") for e in A] + [os.path.join(DL, "L2MN-DEMO_R05_manifest.json")])
        check("the written files pass the format checker (schema and fingerprint)", code == 0 and "warning" not in buf.getvalue(), buf.getvalue())
        check("every assignment file cites the handbook's default rule in its rule header", all(a["rule"]["name"] == "keep_same" and a["rule"]["rotation_share"] == 0.1 and a["rule"]["seed"] == 20261027 for a in A.values()))
        # the enumerator view opens a file the leader view wrote
        b.click('[data-act="view"][data-v="enum"]')
        b.set_files("#f-assign", [os.path.join(DL, "E07_assignment_R05.json")]); b.wait("!!APP.S")
        check("the enumerator view opens a file the leader view wrote, its fingerprint confirmed", b.js("APP.S.events[0].fingerprint_ok") == 1 and b.js("APP.S.assignment.households.length") == len(A["E07"]["households"]))
        b.click('[data-act="close"]'); b.wait("APP.modal"); b.js("(document.querySelector('[data-act=\"modalalt\"]') || document.querySelector('[data-act=\"modalok\"]')).click(); true"); b.wait("APP.S === null")
        b.click('[data-act="view"][data-v="leader"]')
        # protected files (D11)
        tick('[data-lset="protect"]', True)
        before = snapshot()
        b.click('[data-lact="write"]'); b.wait("!!LEAD.ui.passphrases", timeout=60)
        pw = {p["enumerator_id"]: p["passphrase"] for p in b.js("LEAD.ui.passphrases")}
        b.shot(os.path.join(SCREENS, "leader_03_files.png"), full=True)
        b.click('[data-lact="dlall"]')
        f7 = wait_new("E07_assignment_R05*.json", before)
        env = json.load(open(f7, encoding="utf-8"))
        check("protected files: one passphrase per enumerator, made by the tool and shown once; Python opens E07's file with E07's",
              len(set(pw.values())) == 3 and env["schema"] == "rtm-sealed-1.0" and py_unseal(env, pw["E07"])["enumerator_id"] == "E07" and "+000" not in open(f7).read())
        try:
            py_unseal(env, pw["E01"]); wrong_ok = False
        except Exception:
            wrong_ok = True
        check("another enumerator's passphrase does not open the file", wrong_ok)
        check("passphrases are not kept in the leader's session", not any(p in b.js("JSON.stringify(LEAD.state())") for p in pw.values()))

        # ================= B. importing the round's logs =================
        start()
        b.click('[data-lact="tab"][data-t="import"]')
        b.js("RTM.setNow('2026-11-07T09:00:00+08:00'); true")
        b.set_files("#lf-manifest", [MANIFEST]); b.wait("!!LEAD.state().manifest")
        b.set_files("#lf-imports", EXPORTS); b.wait(f"LEAD.state().imports.length === {len(EXPORTS)}")
        b.shot(os.path.join(SCREENS, "leader_04_import.png"), full=True)
        check("a clean batch merges with no question open", b.js("document.querySelector('.panel').innerText").count("No merge question is open") == 1)
        exp_calls = open(os.path.join(S, "expected_R05_calls.csv"), encoding="utf-8").read().strip()
        exp_acts = open(os.path.join(S, "expected_R05_actions.csv"), encoding="utf-8").read().strip()
        b.click('[data-lact="tab"][data-t="calllog"]')
        before = snapshot()
        for w in ("calls", "actions", "interviews", "outcomes", "calendar"):
            b.click(f'[data-lact="export"][data-what="{w}"]')
        fc = wait_new("L2MN-DEMO_R05_calls*.csv", before); fa = wait_new("L2MN-DEMO_R05_actions*.csv", before)
        fi = wait_new("interviews*.csv", before); fo = wait_new("outcome_codes*.csv", before)
        check("the merged call log equals the Python generator's own merge, byte for byte", open(fc, encoding="utf-8").read().strip() == exp_calls)
        check("the merged actions equal the Python merge", open(fa, encoding="utf-8").read().strip() == exp_acts)
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = cf.main([fc])
        check("the merged call log passes the format checker as a call-log CSV", "ERROR" not in buf.getvalue(), buf.getvalue())
        iv = list(csv.DictReader(open(fi, encoding="utf-8")))
        oc = {r["code"] for r in csv.DictReader(open(fo, encoding="utf-8"))}
        man = json.load(open(MANIFEST, encoding="utf-8"))
        check("interviews.csv: one row per household of the round, hhid first, every outcome documented in outcome_codes.csv",
              len(iv) == len({r["hhid"] for r in iv}) and list(iv[0].keys())[0] == "hhid" and {r["outcome"] for r in iv} <= oc and len(iv) >= len(man["households"]))
        # completes per enumerator against an independent count from the merged log
        py_cmp = ref_stats.completes_by_enumerator(exp_calls, man)
        prog = {r["enumerator_id"]: r for r in b.js("LCORE.progress(LEAD.state().manifest, LCORE.cases(LEAD.state().manifest, LCORE.merge(LEAD.state().manifest, LEAD.state().imports.map(f => LCORE.readImport(f.name, f.text)), new Map(), RTM.now())), LCORE.merge(LEAD.state().manifest, LEAD.state().imports.map(f => LCORE.readImport(f.name, f.text)), new Map(), RTM.now())).list")}
        check("progress: completes per enumerator equal an independent count from the merged log", {e: prog[e]["complete"] for e in py_cmp} == py_cmp, [py_cmp, {e: prog[e]["complete"] for e in prog}])
        check("interviews.csv counts the same completes", sum(1 for r in iv if r["outcome"] == "COMPLETE") == sum(py_cmp.values()))
        b.click('[data-lact="tab"][data-t="progress"]')
        b.shot(os.path.join(SCREENS, "leader_05_progress.png"), full=True)
        check("the contacts file's new number is listed for the leader, its digest matching the call log",
              b.js("LCORE.merge(LEAD.state().manifest, LEAD.state().imports.map(f => LCORE.readImport(f.name, f.text)), new Map(), RTM.now()).updates.every(u => u.digest_ok)"))

        # the merge cases (proposal F)
        def issues():
            return b.js("LCORE.merge(LEAD.state().manifest, LEAD.state().imports.map(f => LCORE.readImport(f.name, f.text)), new Map(LEAD.state().resolutions.map(r => [r.key, r])), RTM.now()).issues.map(i => [i.type, i.status, i.key, i.enumerator || '', i.hhid || '', (i.choices || []).map(c => c.value)])")

        b.click('[data-lact="tab"][data-t="import"]')
        b.set_files("#lf-imports", sorted(glob.glob(os.path.join(S, "merge_cases", "duplicate", "*"))))
        b.wait(f"LEAD.state().imports.length === {len(EXPORTS) + 1}")
        check("a file uploaded twice: the copy is ignored and reported", any(i[0] == "duplicate" for i in issues()))
        b.set_files("#lf-imports", sorted(glob.glob(os.path.join(S, "merge_cases", "early_export", "*"))))
        b.wait(f"LEAD.state().imports.length === {len(EXPORTS) + 3}")
        files = b.js("LCORE.merge(LEAD.state().manifest, LEAD.state().imports.map(f => LCORE.readImport(f.name, f.text)), new Map(), RTM.now()).files")
        check("two exports from one enumerator: the later is used, the earlier noted in the files and as a line of the merge report, no gap",
              any("replaced by the later export, E07_2026-11-06.json" in f["use"] for f in files) and any(i[0] == "older_export" and i[3] == "E07" for i in issues())
              and not any(i[0] == "gap" for i in issues()))
        # a later same-day export of E02 whose JSON did not arrive: its CSV is used, as unverified, with the earlier JSON's actions
        n_e02 = b.js("LCORE.merge(LEAD.state().manifest, LEAD.state().imports.map(f => LCORE.readImport(f.name, f.text)), new Map(), RTM.now()).perEnum.get('E02').dials.length")
        b.set_files("#lf-imports", sorted(glob.glob(os.path.join(S, "merge_cases", "csv_only", "*"))))
        b.wait(f"LEAD.state().imports.length === {len(EXPORTS) + 4}")
        g2 = b.js("""(() => { const G = LCORE.merge(LEAD.state().manifest, LEAD.state().imports.map(f => LCORE.readImport(f.name, f.text)), new Map(), RTM.now()), x = G.perEnum.get('E02');
                      return { kind: x.file.kind, unverified: !!x.unverified, dials: x.dials.length, actions: x.actions.length,
                               issues: G.issues.filter(i => i.enumerator === 'E02').map(i => [i.type, i.detail]), uses: G.files.filter(f => f.enumerator_id === 'E02').map(f => [f.name, f.kind, f.use]) }; })()""")
        types2 = [t for t, _ in g2["issues"]]
        check("a later export whose JSON did not arrive: its CSV is used, as unverified, with one more dial; the earlier JSON is noted as replaced (F)",
              g2["kind"] == "csv" and g2["unverified"] and g2["dials"] == n_e02 + 1 and "unverified" in types2 and "older_export" in types2
              and "csv_differs" not in types2 and any(k == "json" and u.startswith("replaced by the later export") for _, k, u in g2["uses"]), g2)
        # review cycle 3: a later CSV lacking an earlier dial; an earlier CSV-only export; a contact update in two files
        rows_l = open(os.path.join(S, "merge_cases", "csv_only", "E02_2026-11-06.csv"), encoding="utf-8").read().rstrip("\n").split("\n")
        hdr = rows_l[0].split(","); ai, hi = hdr.index("action"), hdr.index("hhid")
        rel = next(r for r in rows_l[1:] if r.split(",")[ai] == "RELEASE")
        later = "\n".join([rows_l[0]] + [r for r in rows_l[1:] if r != rel]) + "\n"
        rel_hh, rel_id = rel.split(",")[hi], rel.split(",")[hdr.index("event_id")]
        g4 = b.js(f"""(() => {{ const st = LEAD.state(), seen = new Set(), base = st.imports.filter(f => {{ if (seen.has(f.name) || f.name.includes('(1)')) return false; seen.add(f.name); return true; }}).map(f => LCORE.readImport(f.name, f.text));
                     const imp = base.concat([LCORE.readImport('E02_2026-11-07.csv', {json.dumps(later)})]);
                     const run = res => {{ const G = LCORE.merge(st.manifest, imp, res, RTM.now()), x = G.perEnum.get('E02');
                       return {{ G, x, dial: x.dials.some(d => d.event_id === {json.dumps(rel_id)}), act: x.actions.some(a => a.event_id === {json.dumps(rel_id)} && a.action === 'RELEASE') }}; }};
                     const a = run(new Map()), gap = a.G.issues.find(i => i.type === 'gap' && i.enumerator === 'E02');
                     const b2 = gap ? run(new Map([[gap.key, {{ key: gap.key, decision: 'exclude', logged_by: 'L01', logged_at: '2026-11-07T09:00:00+08:00' }}]])) : null;
                     return {{ file: a.x.file.name, gap: !!gap, keptDial: a.dial, keptAct: a.act, outDial: b2 && b2.dial, outAct: b2 && b2.act }}; }})()""")
        check("a later CSV lacking a dial an earlier export has: the gap is flagged and the dial kept with its release; 'leave them out' drops both",
              g4["file"] == "E02_2026-11-07.csv" and g4["gap"] and g4["keptDial"] and g4["keptAct"] and g4["outDial"] is False and g4["outAct"] is False, g4)
        # review cycle 4: the same later CSV under the same-day name, and an earlier same-day CSV saved as a " (1)" copy
        g6 = b.js(f"""(() => {{ const st = LEAD.state(), seen = new Set(), base = st.imports.filter(f => {{ if (seen.has(f.name) || f.name.includes('(1)')) return false; seen.add(f.name); return true; }}).map(f => LCORE.readImport(f.name, f.text));
                     const G = LCORE.merge(st.manifest, base.concat([LCORE.readImport('E02_2026-11-06.csv', {json.dumps(later)})]), new Map(), RTM.now()), x = G.perEnum.get('E02');
                     return {{ kind: x.file.kind, unverified: !!x.unverified, gap: G.issues.some(i => i.type === 'gap' && i.enumerator === 'E02'), differs: G.issues.some(i => i.type === 'csv_differs' && i.enumerator === 'E02') }}; }})()""")
        check("a later CSV under the same-day name that lacks a dial voided since is used, as unverified, and the gap is flagged (F)",
              g6["kind"] == "csv" and g6["unverified"] and g6["gap"] and not g6["differs"], g6)
        rows_x = open(os.path.join(S, "exports_R05", "E02_2026-11-06.csv"), encoding="utf-8").read().rstrip("\n").split("\n")
        li = hdr.index("logged_at")
        cut = sorted(r.split(",")[li] for r in rows_x[1:])[-6]          # an export made before the last five dials were logged
        earlier_csv = "\n".join([rows_x[0]] + [r for r in rows_x[1:] if r.split(",")[li] <= cut]) + "\n"
        g7 = b.js(f"""(() => {{ const st = LEAD.state(), seen = new Set(), base = st.imports.filter(f => {{ if (seen.has(f.name) || f.name.includes('(1)')) return false; seen.add(f.name); return true; }}).map(f => LCORE.readImport(f.name, f.text));
                     const G = LCORE.merge(st.manifest, base.concat([LCORE.readImport('E02_2026-11-06 (1).csv', {json.dumps(earlier_csv)})]), new Map(), RTM.now());
                     const row = G.files.find(f => f.name === 'E02_2026-11-06 (1).csv');
                     return {{ used: G.perEnum.get('E02').file.name, use: row && row.use, differs: G.issues.some(i => i.type === 'csv_differs'), older: G.issues.some(i => i.type === 'older_export' && i.file === 'E02_2026-11-06 (1).csv') }}; }})()""")
        check("an earlier same-day CSV saved as a ' (1)' copy is an earlier export: replaced by the later one and noted, not flagged as differing",
              g7["used"] == "E02_2026-11-06.json" and not g7["differs"] and g7["older"] and (g7["use"] or "").startswith("replaced by the later export"), g7)
        mid = sorted(rows_x[1:], key=lambda r: r.split(",")[li])[len(rows_x) // 2]       # a dial logged mid-day, voided before a re-export
        voided_csv = "\n".join([rows_x[0]] + [r for r in rows_x[1:] if r != mid]) + "\n"
        g8 = b.js(f"""(() => {{ const st = LEAD.state(), seen = new Set(), base = st.imports.filter(f => {{ if (seen.has(f.name) || f.name.includes('(1)')) return false; seen.add(f.name); return true; }}).map(f => LCORE.readImport(f.name, f.text));
                     const extra = LCORE.readImport('E02_2026-11-06 (2).csv', {json.dumps(voided_csv)});
                     const G = LCORE.merge(st.manifest, base.concat([extra]), new Map(), RTM.now());
                     const last = {json.dumps(max(r.split(",")[li] for r in rows_x[1:]))};
                     return {{ differs: G.issues.some(i => i.type === 'csv_differs' && i.file === 'E02_2026-11-06 (2).csv'), stamp: extra.exported_at === last }}; }})()""")
        check("a same-day CSV lacking a dial logged before its own time is flagged as differing, not taken for an earlier export; a CSV's time stays in the enumerator's clock",
              g8["differs"] and g8["stamp"], g8)
        rows_e = open(os.path.join(S, "merge_cases", "early_export", "E07_2026-10-30.csv"), encoding="utf-8").read().rstrip("\n").split("\n")
        fake = rows_e[-1].split(","); fake[hdr.index("event_id")] = "00000000-0000-4000-8000-000000000001"
        early = "\n".join(rows_e + [",".join(fake)]) + "\n"
        g5 = b.js(f"""(() => {{ const st = LEAD.state(), imp = st.imports.filter(f => f.name.startsWith('E0') && !f.name.includes('(1)')).map(f => LCORE.readImport(f.name, f.text))
                       .filter(f => !f.name.startsWith('E07_2026-10-30')).concat([LCORE.readImport('E07_2026-10-30.csv', {json.dumps(early)})]);
                     const G = LCORE.merge(st.manifest, imp, new Map(), RTM.now()), x = G.perEnum.get('E07');
                     return {{ file: x.file.name, gap: G.issues.some(i => i.type === 'gap' && i.file === 'E07_2026-10-30.csv'), kept: x.dials.some(d => d.event_id === '00000000-0000-4000-8000-000000000001') }}; }})()""")
        check("an earlier export that came as a CSV only is compared too: a dial the later JSON lacks is kept and the gap flagged (F)",
              g5["file"] == "E07_2026-11-06.json" and g5["gap"] and g5["kept"], g5)
        b.click('[data-lact="clearimports"]'); b.wait("LEAD.state().imports.length === 0")
        b.set_files("#lf-imports", EXPORTS + sorted(glob.glob(os.path.join(S, "merge_cases", "duplicate", "*"))) + sorted(glob.glob(os.path.join(S, "merge_cases", "early_export", "*"))))
        b.wait(f"LEAD.state().imports.length === {len(EXPORTS) + 3}")
        b.set_files("#lf-imports", sorted(glob.glob(os.path.join(S, "merge_cases", "wrong_round", "*"))))
        b.wait(f"LEAD.state().imports.length === {len(EXPORTS) + 4}")
        check("a file from another round is not merged, and is reported", any(i[0] == "wrong_round" for i in issues()))
        cf = json.load(open(os.path.join(S, "merge_cases", "conflict", "E01_2026-11-06_contacts.json"), encoding="utf-8")); cf["round"] = 4
        g3 = b.js(f"""(() => {{ const st = LEAD.state(), extra = LCORE.readImport('E01_2026-10-21_contacts.json', {json.dumps(json.dumps(cf))});
                     const G = LCORE.merge(st.manifest, st.imports.map(f => LCORE.readImport(f.name, f.text)).concat([extra]), new Map(), RTM.now());
                     return {{ wrong: G.issues.some(i => i.type === 'wrong_round' && i.file === 'E01_2026-10-21_contacts.json'), use: G.files.find(f => f.name === 'E01_2026-10-21_contacts.json').use }}; }})()""")
        check("a contacts file from another round is not used, and is reported", g3["wrong"] and g3["use"] == "other round, not merged", g3)
        b.set_files("#lf-imports", sorted(glob.glob(os.path.join(S, "merge_cases", "conflict", "*"))))
        b.wait("LEAD.state().imports.some(f => f.name.includes('conflict') || f.text.includes('18:55'))", timeout=10)
        time.sleep(0.5)
        n_files = b.js("LEAD.state().imports.filter(f => f.name.endsWith('_contacts.json')).length")
        n_upd = b.js("LCORE.merge(LEAD.state().manifest, LEAD.state().imports.map(f => LCORE.readImport(f.name, f.text)), new Map(), RTM.now()).updates.length")
        n_ev = len({u["event_id"] for f in [os.path.join(S, "exports_R05", "E01_2026-11-06_contacts.json"), os.path.join(S, "merge_cases", "conflict", "E01_2026-11-06_contacts.json")]
                    for u in json.load(open(f, encoding="utf-8"))["updates"]})
        check("a contact update repeated in a later contacts file counts once", n_files == 2 and n_upd == n_ev == 1, [n_files, n_upd, n_ev])
        iss = issues()
        tc = [i for i in iss if i[0] == "two_completes"]
        na = [i for i in iss if i[0] == "not_assigned"]
        check("dials on a household assigned to someone else are kept and flagged", na and na[0][1] == "open", na)
        check("same-day exports that share a name are each checked against their own JSON (no false 'CSV differs')", not any(i[0] == "csv_differs" for i in iss), [i for i in iss if i[0] == "csv_differs"])
        check("two completes for one household are kept, flagged, and left out of the counts until the leader chooses", tc and tc[0][1] == "open"
              and b.js(f"LCORE.cases(LEAD.state().manifest, LCORE.merge(LEAD.state().manifest, LEAD.state().imports.map(f => LCORE.readImport(f.name, f.text)), new Map(), RTM.now())).get('{tc[0][4]}').state") == "conflict", tc)
        b.js("APP.render(); true")
        b.shot(os.path.join(SCREENS, "leader_06_decisions.png"), full=True)
        # the leader decides: a mistyped ID
        b.click(f'[data-lact="resolve"][data-key="{na[0][2]}"][data-v="typo"]')
        b.wait("LEAD.state().resolutions.length === 1")
        iss = issues()
        check("a decision is logged with who and when; the mistyped dials leave the counts and the conflict closes",
              b.js("LEAD.state().resolutions[0].logged_by") == "L01" and not any(i[0] == "two_completes" for i in iss)
              and b.js(f"LCORE.cases(LEAD.state().manifest, LCORE.merge(LEAD.state().manifest, LEAD.state().imports.map(f => LCORE.readImport(f.name, f.text)), new Map(LEAD.state().resolutions.map(r => [r.key, r])), RTM.now())).get('{tc[0][4]}').state") == "complete")
        # an edited file
        b.set_files("#lf-imports", sorted(glob.glob(os.path.join(S, "merge_cases", "edited", "*"))))
        time.sleep(0.8)
        iss = issues()
        check("a file edited after export: its chain fails, the differing event is a conflict", any(i[0] == "chain_broken" for i in iss) and any(i[0] == "event_conflict" for i in iss))
        ec = next(i for i in iss if i[0] == "event_conflict")
        b.click(f'[data-lact="resolve"][data-key="{ec[2]}"][data-v="{ec[5][0]}"]')
        b.wait("LEAD.state().resolutions.length === 2")
        check("the leader keeps the verified version: the event counts again", not any(i[0] == "event_conflict" and i[1] == "open" for i in issues()))
        before = snapshot()
        b.click('[data-lact="export"][data-what="report"]')
        fr = wait_new("L2MN-DEMO_R05_merge_report*.csv", before)
        rep = list(csv.DictReader(open(fr, encoding="utf-8")))
        check("the merge report lists every situation, the rule applied and the leader's decisions",
              {"duplicate", "wrong_round", "not_assigned", "chain_broken", "event_conflict"} <= {r["type"] for r in rep} and any(r["decision"] == "typo" and r["resolved_by"] == "L01" for r in rep))
        # save and reopen the leader's session
        st_before = b.js("JSON.stringify({ i: LEAD.state().imports.length, r: LEAD.state().resolutions, m: LEAD.state().manifest.fingerprint })")
        before = snapshot()
        b.click('[data-lact="save"]'); b.wait("APP.modal"); b.click('[data-act="modalok"]')
        fp = wait_new("leader_L2MN-DEMO_R05*.json", before)
        P = json.load(open(fp, encoding="utf-8")); chk = P.pop("checksum")
        check("the leader session saves with a checksum that verifies in Python", sha(canon(P)) == chk)
        b.click('[data-lact="close"]'); b.wait("APP.modal"); b.js("(document.querySelector('[data-act=\"modalalt\"]') || document.querySelector('[data-act=\"modalok\"]')).click(); true"); b.wait("!LEAD.state()")
        b.set_files("#lf-project", [fp]); b.wait("!!LEAD.state()")
        check("reopening the leader session restores imports, decisions and the manifest", b.js("JSON.stringify({ i: LEAD.state().imports.length, r: LEAD.state().resolutions, m: LEAD.state().manifest.fingerprint })") == st_before)
        bad = os.path.join(SCRATCH, f"leader_tampered_{STAMP}.json")
        T = json.load(open(fp, encoding="utf-8")); T["resolutions"][0]["decision"] = "reassign"; json.dump(T, open(bad, "w"))
        b.click('[data-lact="close"]'); b.wait("APP.modal"); b.js("(document.querySelector('[data-act=\"modalalt\"]') || document.querySelector('[data-act=\"modalok\"]')).click(); true"); b.wait("!LEAD.state()")
        b.set_files("#lf-project", [bad]); time.sleep(1)
        check("a leader session changed outside the tool is refused", not b.js("!!LEAD.state()") and b.js("APP.toasts.some(t => /checksum/.test(t.text))"))
        # protected leader session
        b.set_files("#lf-project", [fp]); b.wait("!!LEAD.state()")
        before = snapshot()
        b.click('[data-lact="save"]'); b.wait("APP.modal")
        b.js("document.querySelector('#modal-input').value = 'leader-secret-1'; document.querySelector('#modal-input2').value = 'leader-secret-1'; true")
        b.click('[data-act="modalok"]')
        fps = wait_new("leader_L2MN-DEMO_R05*.json", before)
        envs = json.load(open(fps, encoding="utf-8"))
        check("a protected leader session holds no phone number in the clear; Python opens it", envs["schema"] == "rtm-sealed-1.0" and "+000" not in open(fps).read()
              and py_unseal(envs, "leader-secret-1")["schema"] == "rtm-leader-1.0")

        # ================= C. analytics, against an independent implementation =================
        start()
        b.set_files("#lf-hist", HIST); b.wait(f"LEAD.state().inputs.history.length === {len(HIST)}")
        b.click('[data-lact="tab"][data-t="import"]')
        b.set_files("#lf-manifest", [MANIFEST]); b.wait("!!LEAD.state().manifest")
        b.set_files("#lf-imports", EXPORTS); b.wait(f"LEAD.state().imports.length === {len(EXPORTS)}")
        b.click('[data-lact="tab"][data-t="analytics"]')
        b.shot(os.path.join(SCREENS, "leader_07_analytics.png"), full=True)
        js = b.js("""(() => { const d = LCORE.buildHistory(LEAD.state().inputs.history.map(f => LCORE.readHistoryFile(f.name, f.text)), {tz: 'Asia/Ulaanbaatar'});
            const input = []; for (const [r, ds] of d.dials) ds.forEach(x => input.push({hhid: x.hhid, round: r, wall: x.wall, status: x.status, cmp_code: x.cmp, dur: x.dur, operator: x.operator}));
            const T = LSTATS.build(input), rs = T.rounds, I = LSTATS.interviewed(T, null);
            const s33 = LSTATS.s33(T, rs), s34 = LSTATS.s34(I, rs, null), s31 = LSTATS.s31(T, rs), s35 = LSTATS.s35(T, rs, {});
            return JSON.stringify({ cases: T.cases.size, dials: T.dials.length, fallbacks: T.fallbacks, dpc: s33.pooled.dials_per_complete, byRound: s33.byRound.map(x => [x.round, x.cases, x.completes, x.dials]),
              hazard: s33.hazard.slice(0, 8).map(h => [h.k, h.at_risk, h.completes]), cont: s34.continuation.map(x => [x.from, x.round, x.base, Math.round(x.share * 1e6)]),
              missed: [s33.missed.newly_missed, s33.missed.redialled && s33.missed.redialled.k],
              cpd: Math.round(s31.pooled.completes_per_day * 1e6), od: s31.pooled.enumerator_days, att: Math.round(s31.pooled.attendance * 1e6), mpc: Math.round(s31.pooled.minutes_per_complete * 1e4),
              hours: s35.hours.map(h => [h.hour, h.n, h.success ? h.success.k : null]), same: s35.same_time.map(x => x.n), cont35: [s35.continuity.n, Math.round(s35.continuity.same_share * 1e6)] }); })()""")
        js = json.loads(js)
        py = ref_stats.reference(HIST)
        for k in ("cases", "dials", "fallbacks", "byRound", "hazard", "cont", "missed", "od", "cpd", "att", "mpc", "hours", "same", "cont35"):
            check(f"analytics '{k}' equals the independent Python implementation of the handbook's rules", js[k] == py[k], [js[k], py[k]])
        check("dials per complete, pooled, equals the reference", abs(js["dpc"] - py["dpc"]) < 1e-9, [js["dpc"], py["dpc"]])
        txt = b.js("document.querySelector('.panel').innerText")
        check("the analytics tab shows each measure with its handbook definition", all(s in txt for s in ("Box 2.0", "Section 3.3.1", "Section 3.1", "Section 3.5.2", "Section 3.5.3", "Section 3.5.4")))

        # earlier rounds given as interview records only (a program without dial logs)
        ivp = os.path.join(SCRATCH, f"interviews_R01_R04_{STAMP}.csv")
        done = {}
        for pth in [x for x in HIST if x.endswith("_calls.csv")]:
            for r in csv.DictReader(open(pth, encoding="utf-8")):
                if r["status"] == "COMPLETE" and (r["dur_min"] == "" or float(r["dur_min"]) >= 3):
                    done.setdefault((r["hhid"], r["round"]), (r["ts"][:10], r["operator"]))
        with open(ivp, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f, lineterminator="\n"); w.writerow(["hhid", "round", "interview_date", "outcome", "enumerator"])
            for (h, rnd), (d, e) in sorted(done.items()):
                w.writerow([h, rnd, d, "complete", e])
        start()
        b.set_files("#lf-hist", [ivp]); b.wait("LEAD.state().inputs.history.length === 1")
        js_iv = b.js("[...LCORE.buildHistory(LEAD.state().inputs.history.map(f => LCORE.readHistoryFile(f.name, f.text)), {tz: 'Asia/Ulaanbaatar'}).complete].map(([r, m]) => [r, m.size])")
        check("earlier rounds given as interview records only are read: interviewed households per round as from the dial logs",
              [list(x) for x in js_iv] == py_int, [js_iv, py_int])

        # an Excel workbook as the household list (proposal G: a built-in reader)
        try:
            import openpyxl
        except ImportError:
            openpyxl = None
        if openpyxl:
            hl_rows = list(csv.reader(open(os.path.join(S, "household_list.csv"), encoding="utf-8")))
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "households"
            for r in hl_rows:
                ws.append(r)
            for row in ws.iter_rows():
                for c in row:
                    c.number_format = "@"
            x1 = os.path.join(SCRATCH, f"household_list_{STAMP}.xlsx"); wb.save(x1)
            wb2 = openpyxl.Workbook(); ws2 = wb2.active
            for i, r in enumerate(hl_rows):
                ws2.append([int(v.replace("+", "").replace(" ", "")) if (i and j in (4, 6, 8) and v) else v for j, v in enumerate(r)])
            x2 = os.path.join(SCRATCH, f"household_list_numbers_{STAMP}.xlsx"); wb2.save(x2)
            start()
            b.set_files("#lf-hh", [x1]); b.wait("!!LEAD.state().inputs.households")
            same = b.js(f"JSON.stringify([...LCORE.parseHouseholds(LEAD.state().inputs.households.text).households]) === JSON.stringify([...LCORE.parseHouseholds({json.dumps(open(os.path.join(S, 'household_list.csv'), encoding='utf-8').read())}).households])")
            check("an Excel workbook is read offline by the built-in reader, exactly as the same list in CSV", same and "110 households to assign" in b.js("document.querySelector('.panel').innerText"))
            b.set_files("#lf-hh", [x2]); b.wait("APP.toasts.some(t => /stored as numbers/.test(t.text))")
            check("phone numbers stored as numbers in Excel raise a warning: leading zeros and + are lost", "phone_1" in b.js("APP.toasts.map(t => t.text).join(' ')"))
        else:
            check("an Excel workbook is read by the built-in reader (openpyxl missing: not tested)", False, "install openpyxl to run this check")

        # ================= D. matching criteria, the day list, and moving households during the round =================
        # (the author's requests of 28 October 2026: assign by language, sex or dialect and record the criteria; longer day
        # lists so that nobody waits; move waiting households from slower to faster enumerators without double calls)
        import check_formats as cfm
        b.ws.call("Emulation.clearDeviceMetricsOverride")
        b.js("RTM.setNow('2026-10-23T15:00:00+08:00'); true")
        start()
        b.set_files("#lf-hh", [os.path.join(S, "household_list.csv")]); b.wait("!!LEAD.state().inputs.households")
        b.set_files("#lf-hist", HIST); b.wait(f"LEAD.state().inputs.history.length === {len(HIST)}")
        b.set_files("#lf-repl", [os.path.join(S, "selection.csv")]); b.wait("!!LEAD.state().inputs.replacements")
        setv('[data-lset="cal.window_start"]', "2026-10-26"); setv('[data-lset="cal.window_end"]', "2026-11-20")
        tick('[data-lteam="2"][data-f="available"]', False)
        setv('[data-lset="rule.rotation_share"]', "10"); setv('[data-lset="rule.seed"]', "20261027")
        txt = b.js("document.querySelector('.panel').innerText")
        check("step 5 offers the household list's columns as matching criteria, with their values (language: kk, mn; respondent sex: f, m)",
              "Match households to enumerators by" in txt and "language" in txt and "respondent sex" in txt and "(f, m;" in txt and "(kk, mn;" in txt, txt[:300])
        tick('[data-lcrit="language"]', True)
        b.click('[data-lact="apply"]')
        check("with a criterion ticked, the rule refuses to run until every available enumerator has values for it", any("values for language" in e for e in b.js("LEAD.ui.errors")), b.js("LEAD.ui.errors"))
        team = b.js("LEAD.state().team.map(t => t.id)")
        for e, v in {"E01": "mn kk", "E02": "mn", "E07": "mn"}.items():
            setv(f'[data-lattr="{team.index(e)}"][data-c="language"]', v)
        b.click('[data-lact="apply"]'); b.wait("!!LEAD.state().plan")
        b.shot(os.path.join(SCREENS, "leader_09_criteria.png"), full=True)
        kk = {r["hhid"] for r in hl if r["language"] == "kk"}
        out = {h: x for h, x in json.loads(b.js("JSON.stringify([...LEAD.state().plan.out])"))}
        kk_in = sorted(h for h in kk if h in out)
        check("every Kazakh-speaking household is with the one enumerator who speaks Kazakh, and no household is left without an enumerator",
              kk_in and all(out[h]["enumerator_id"] == "E01" for h in kk_in) and len(out) == 110 and b.js("LEAD.state().plan.unmatched.length") == 0, {h: out[h]["enumerator_id"] for h in kk_in})
        moved = [h for h, x in out.items() if x["reason"] == "match"]
        check("a kept household whose enumerator does not match moves with the reason 'match' and its previous enumerator recorded",
              moved and all(h in kk and out[h]["previous_enumerator"] in ("E02", "E07") for h in moved), moved)
        check("every household is marked as matched, and the note counts the moves", all(x.get("matched") is True for x in out.values()) and any("moved to match the criteria" in n for n in b.js("LEAD.state().plan.notes")), b.js("LEAD.state().plan.notes"))
        if b.js("!!document.querySelector('[data-lact=\"balance\"]')"):
            b.click('[data-lact="balance"]'); time.sleep(0.3)
            out2 = {h: x for h, x in json.loads(b.js("JSON.stringify([...LEAD.state().plan.out])"))}
            check("balancing the loads respects the criteria: no Kazakh-speaking household leaves E01", all(out2[h]["enumerator_id"] == "E01" for h in kk_in))
        else:
            check("balancing the loads respects the criteria (loads already within the margin: not exercised)", True)
        # nobody speaks Kazakh: the households wait for the leader
        setv(f'[data-lattr="{team.index("E01")}"][data-c="language"]', "mn")
        b.click('[data-lact="apply"]'); b.wait("!!LEAD.state().plan")
        un = b.js("LEAD.state().plan.unmatched")
        check("when no available enumerator matches, the households are listed without an enumerator and no file can be written",
              sorted(un) == kk_in and b.js("document.querySelector('[data-lact=\"write\"]').disabled") and any("match no available enumerator" in w for w in b.js("LEAD.state().plan.warnings")), un)
        setv(f'[data-lmove="{un[0]}"]', "E02")
        check("the leader assigns such a household by hand: it leaves the list, marked as not matched, with a warning",
              b.js("LEAD.state().plan.unmatched.length") == len(un) - 1 and b.js(f"LEAD.state().plan.out.get('{un[0]}').matched") is False and b.js("APP.toasts.some(t => /does not match E02/.test(t.text))"))
        setv(f'[data-lattr="{team.index("E01")}"][data-c="language"]', "kk, mn")
        b.click('[data-lact="apply"]'); b.wait("!!LEAD.state().plan")
        check("the day list is taken from the last round's manifest, like the dial cap and the call hours", b.js("LEAD.state().cal.day_list") == 5)
        setv('[data-lset="cal.day_list"]', "")
        dl_default = b.js("LEAD.state().plan.dayListDefault")
        check("left empty, the day list defaults to 1.5 × the busiest enumerator's planned first dials per day, shown in step 4",
              dl_default >= 1 and f"empty means 1.5 × that, {dl_default}" in b.js("document.querySelector('.panel').innerText") and b.js("LEAD.state().cal.day_list") is None, dl_default)
        before = snapshot()
        b.click('[data-lact="write"]'); b.wait("!!LEAD.ui.ready", timeout=30)
        b.click('[data-lact="dlall"]')
        got = [wait_new(n, before) for n in ("E01_assignment_R05.json", "E02_assignment_R05.json", "E07_assignment_R05.json", "L2MN-DEMO_R05_manifest.json")]
        A = {e: json.load(open(os.path.join(DL, f"{e}_assignment_R05.json"), encoding="utf-8")) for e in ("E01", "E02", "E07")}
        M = json.load(open(os.path.join(DL, "L2MN-DEMO_R05_manifest.json"), encoding="utf-8"))
        check("the manifest records the criteria and every enumerator's values; each file carries the criteria and its enumerator's values",
              all(got) and M.get("matching") == {"criteria": ["language"], "enumerators": {"E01": {"language": ["kk", "mn"]}, "E02": {"language": ["mn"]}, "E07": {"language": ["mn"]}}}
              and all(a["matching"] == {"criteria": ["language"], "enumerator": M["matching"]["enumerators"][e]} for e, a in A.items()), M.get("matching"))
        check("each household in a file carries its own value and whether its enumerator matches; the kk households are all in E01's file",
              all(h["criteria"]["language"] in ("kk", "mn") and h["assignment"]["matched"] is True for a in A.values() for h in a["households"])
              and sorted(h["hhid"] for h in A["E01"]["households"] if h["criteria"]["language"] == "kk") == kk_in
              and not any(h["criteria"]["language"] == "kk" for e in ("E02", "E07") for h in A[e]["households"]))
        check("the files carry the day list, equal to the default when the leader left it empty", all(a["calendar"]["day_list"] == dl_default for a in A.values()) and M["calendar"]["day_list"] == dl_default and all(a["revision"] == 1 for a in A.values()),
              [dl_default, M["calendar"].get("day_list"), {e: (a["calendar"].get("day_list"), a.get("revision")) for e, a in A.items()}])
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = cfm.main([os.path.join(DL, f"{e}_assignment_R05.json") for e in A] + [os.path.join(DL, "L2MN-DEMO_R05_manifest.json")])
        check("files with criteria and a day list pass the format checker", code == 0 and "warning" not in buf.getvalue(), buf.getvalue())
        setv('[data-lset="cal.day_list"]', "9")
        check("changing the day list sets the files aside but keeps the assignment", b.js("LEAD.ui.ready === null && !!LEAD.state().plan"))
        before = snapshot()
        b.click('[data-lact="write"]'); b.wait("!!LEAD.ui.ready", timeout=30)
        b.click('[data-lact="dlall"]')
        wait_new("E07_assignment_R05.json", before)
        check("the leader's own day list is written instead", json.load(open(os.path.join(DL, "E07_assignment_R05.json"), encoding="utf-8"))["calendar"]["day_list"] == 9)
        # the saved session keeps the criteria and the values
        before = snapshot()
        b.click('[data-lact="save"]'); b.wait("APP.modal"); b.click('[data-act="modalok"]')
        fpc = wait_new("leader_L2MN-DEMO_R05*.json", before)
        Pc = json.load(open(fpc, encoding="utf-8"))
        check("the saved leader session keeps the criteria and each enumerator's values", Pc["settings"]["criteria"] == ["language"] and Pc["settings"]["team"][team.index("E01")]["attrs"] == {"language": "kk, mn"})

        # the proposal of moves, on fabricated paces: from the enumerator furthest behind to the one with most spare days
        unit = b.js("""(() => { const rows = [
            { enumerator_id: 'A', known: true, fresh: true, need: 3.4 }, { enumerator_id: 'B', known: true, fresh: true, need: -2.6 },
            { enumerator_id: 'C', known: true, fresh: true, need: -0.4 }, { enumerator_id: 'D', known: true, fresh: false, export_day: '2026-11-01', need: -9, why: 'last export 2026-11-01, not today' } ];
          const C = new Map(); const add = (h, e, state, extra) => C.set(h, Object.assign({ hhid: h, enumerator_id: e, state, dials: [], n: 0 }, extra || {}));
          add('a1', 'A', 'not_started'); add('a2', 'A', 'not_started'); add('a3', 'A', 'in_progress', { dials: [{ status: 'NO_ANSWER' }], n: 1 }); add('a4', 'A', 'complete');
          add('a5', 'A', 'in_progress', { dials: [{ status: 'NO_ANSWER', appointment_ts: '2026-11-10T10:00:00+08:00' }], n: 1 }); add('b1', 'B', 'not_started'); add('d1', 'D', 'not_started');
          const r1 = LCORE.proposeMoves({ rows }, C, { include: 'not_started' });
          const r2 = LCORE.proposeMoves({ rows }, C, { include: 'in_progress' });
          const r3 = LCORE.proposeMoves({ rows }, C, { include: 'not_started', eligible: (h, e) => !(h.hhid === 'a2' && e === 'B'), hh: new Map([['a1', { hhid: 'a1' }], ['a2', { hhid: 'a2' }]]) });
          const r4 = LCORE.proposeMoves({ rows }, C, { include: 'not_started', allowStale: true });
          const mv = (r) => r.moves.map(m => [m.hhid, m.from, m.to]);
          return { r1: mv(r1), r2: mv(r2), r3: mv(r3), r4: mv(r4), left: r1.left_out }; })()""")
        check("moves go one at a time from the enumerator furthest behind to the one with most spare days, never dialed households first, until nobody is more than half a household off",
              unit["r1"] == [["a2", "A", "B"], ["a1", "A", "B"]], unit)
        check("called households without a booked callback move only when the leader allows it; a booked callback or a complete never moves", unit["r2"] == [["a2", "A", "B"], ["a1", "A", "B"], ["a3", "A", "B"]], unit)
        check("a household moves only to an enumerator who matches its criteria", unit["r3"] == [["a1", "A", "B"]], unit)
        check("an enumerator whose export is from an earlier day is not judged unless the leader says so", unit["left"] == [{"id": "D", "why": "last export 2026-11-01, not today"}] and unit["r4"][0] == ["a2", "A", "D"], unit)

        # moving households during the round: the sample's exports of 6 November, judged on Friday 13 November with 7 working days left
        start()
        b.set_files("#lf-hh", [os.path.join(S, "household_list.csv")]); b.wait("!!LEAD.state().inputs.households")
        b.set_files("#lf-hist", HIST); b.wait(f"LEAD.state().inputs.history.length === {len(HIST)}")
        b.click('[data-lact="tab"][data-t="import"]')
        b.js("RTM.setNow('2026-11-13T09:00:00+08:00'); true")
        b.set_files("#lf-manifest", [MANIFEST]); b.wait("!!LEAD.state().manifest")
        b.set_files("#lf-imports", EXPORTS); b.wait(f"LEAD.state().imports.length === {len(EXPORTS)}")
        b.click('[data-lact="tab"][data-t="progress"]')
        txt = b.js("document.querySelector('.panel').innerText")
        check("the Progress tab shows each enumerator's pace and capacity for the working days left, and who is short or has spare days",
              "Move waiting households" in txt and "Completes per day" in txt and "7 working days left" in txt, txt[:200])
        man = json.load(open(MANIFEST, encoding="utf-8"))
        hold = {h["hhid"]: h["enumerator_id"] for h in man["households"]}
        py_pace = ref_stats.pace_by_enumerator(exp_calls, exp_acts, man, days_left=7)
        js_pace = {r["enumerator_id"]: r for r in b.js("LCORE.pace(LEAD.state().manifest, LEAD.merged().C, LEAD.merged().P, RTM.now()).rows")}
        check("pace, capacity and the shortfall per enumerator equal an independent computation from the merged log",
              all(abs(js_pace[e]["pace"] - p["pace"]) < 1e-9 and js_pace[e]["open"] == p["open"] and abs(js_pace[e]["need"] - p["need"]) < 1e-9 for e, p in py_pace.items()), [py_pace, {e: (r["pace"], r["open"], r["need"]) for e, r in js_pace.items()}])
        b.click('[data-lact="propose"]'); time.sleep(0.3)
        check("exports from an earlier day are not judged unless the leader says so", b.js("LEAD.ui.moves.list.length") == 0 and "not today" in b.js("document.querySelector('.panel').innerText"))
        tick('[data-lset="__movesStale"]', True)
        b.click('[data-lact="propose"]'); time.sleep(0.3)
        mv = b.js("LEAD.ui.moves.list")
        waiting = {e: b.js(f"[...LEAD.merged().C.values()].filter(c => c.enumerator_id === '{e}' && c.state === 'not_started').map(c => c.hhid).sort().reverse()") for e in ("E01", "E07")}
        exp_need = {e: round(p["need"], 1) for e, p in py_pace.items()}
        check("on 13 November two households move to E02, the only enumerator with spare days (2.3): first from E01 (1.9 short), then from E07 (1.0 short, now the furthest behind), each its latest-planned household, until E02 is within half a household",
              exp_need == {"E01": 1.9, "E02": -2.3, "E07": 1.0} and len(mv) == 2 and [(m["from"], m["to"], m["state"], m["hhid"]) for m in mv] == [("E01", "E02", "not_started", waiting["E01"][0]), ("E07", "E02", "not_started", waiting["E07"][0])], [mv, exp_need, waiting])
        by_hand = next(h for h in reversed(waiting["E01"]) if h != mv[0]["hhid"])
        b.js(f"document.querySelector('#lmv-hh').value = {json.dumps(by_hand)}; document.querySelector('#lmv-to').value = 'E02'; true")
        b.click('[data-lact="addmove"]'); time.sleep(0.3)
        n_list = b.js("LEAD.ui.moves.list.length")
        b.js("document.querySelector('#lmv-hh').value = [...LEAD.merged().C.values()].find(c => c.state === 'complete').hhid; true"); b.click('[data-lact="addmove"]'); time.sleep(0.2)
        check("the leader adds a move by hand; a household already complete is refused",
              n_list == 3 and b.js("LEAD.ui.moves.list[2].by_hand") is True and b.js("APP.toasts.some(t => /cannot move: it is Complete/.test(t.text)) && LEAD.ui.moves.list.length === 3"),
              [n_list, b.js("APP.toasts.map(t => t.text)")])
        b.shot(os.path.join(SCREENS, "leader_10_moves.png"), full=True)
        old_fp = {e["enumerator_id"]: e["fingerprint"] for e in man["enumerators"]}
        before = snapshot()
        b.click('[data-lact="writemoves"]'); b.wait("!!LEAD.ui.ready && LEAD.ui.readyKind === 'moves'", timeout=30)
        b.click('[data-lact="dlall"]')
        got = [wait_new(n, before) for n in ("E01_assignment_R05.json", "E02_assignment_R05.json", "E07_assignment_R05.json", "L2MN-DEMO_R05_manifest.json")]
        R2 = {e: json.load(open(os.path.join(DL, f"{e}_assignment_R05.json"), encoding="utf-8")) for e in ("E01", "E02", "E07")}
        M2 = json.load(open(os.path.join(DL, "L2MN-DEMO_R05_manifest.json"), encoding="utf-8"))
        moved_ids = {mv[0]["hhid"], mv[1]["hhid"], by_hand}
        check("revised files are written for the three enumerators concerned, and the manifest's revision 2 supersedes revision 1",
              all(got) and M2["revision"] == 2 and M2["supersedes"] == man["fingerprint"] and all(a["revision"] == 2 and a["supersedes"] == old_fp[e] for e, a in R2.items()), [M2.get("revision"), M2.get("supersedes") == man["fingerprint"]])
        check("the manifest lists the three moves with when and by whom, the households now with E02 for the reason 'mid_round', their previous enumerator kept",
              [(m["hhid"], m["from"], m["to"]) for m in M2["moves"]] == [(mv[0]["hhid"], "E01", "E02"), (mv[1]["hhid"], "E07", "E02"), (by_hand, "E01", "E02")] and all(m["by"] == "L01" and m["at"].startswith("2026-11-13T09:00") for m in M2["moves"])
              and all(h["enumerator_id"] == "E02" and h["reason"] == "mid_round" and h["previous_enumerator"] == hold[h["hhid"]] for h in M2["households"] if h["hhid"] in moved_ids), M2["moves"])
        check("the roster counts follow the moves, and each revised file's earlier fingerprint is kept so that logs begun from it still belong",
              {e["enumerator_id"]: e["households"] for e in M2["enumerators"]} == {"E01": 38, "E02": 27, "E07": 45} and all(e["previous_fingerprints"] == [old_fp[e["enumerator_id"]]] for e in M2["enumerators"]),
              {e["enumerator_id"]: e["households"] for e in M2["enumerators"]})
        e02new = {h["hhid"]: h for h in R2["E02"]["households"]}
        check("E02's revised file holds its 24 households and the 3 moved ones, with the previous enumerator's dials (handover), a first dial today and the household's usual hour",
              len(e02new) == 27 and all(e02new[h]["handover"]["dials"] == 0 and e02new[h]["handover"]["from"] == hold[h] and e02new[h]["plan"]["first_dial_date"] == "2026-11-13" and e02new[h]["assignment"]["reason"] == "mid_round" for h in moved_ids)
              and all("handover" not in h for k, h in e02new.items() if k not in moved_ids), [e02new[h].get("handover") for h in moved_ids])
        hl_by = {r["hhid"]: r for r in hl}
        corrected = {}                      # numbers enumerators corrected this round go into the revised files (the contacts file of the exports)
        for cfile in glob.glob(os.path.join(S, "exports_R05", "*_contacts.json")):
            for u in json.load(open(cfile, encoding="utf-8"))["updates"]:
                if u.get("number"): corrected.setdefault(u["hhid"], set()).add(u["number"])
        own = lambda h: {hl_by[h][k] for k in ("phone_1", "phone_2", "phone_3")} | corrected.get(h, set())
        check("access control holds: each revised file carries only its own households' numbers, with the numbers corrected this round",
              all(all(p["number"] in own(h["hhid"]) for h in a["households"] for p in h["contact"]["phones"]) for a in R2.values())
              and any(p["number"] in corrected.get(h["hhid"], set()) for a in R2.values() for h in a["households"] for p in h["contact"]["phones"]), corrected)
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = cfm.main([os.path.join(DL, f"{e}_assignment_R05.json") for e in R2] + [os.path.join(DL, "L2MN-DEMO_R05_manifest.json")])
        check("the revised files and the manifest pass the format checker", code == 0 and "warning" not in buf.getvalue(), buf.getvalue())
        # the merge with the revised manifest: dials before the move are the previous enumerator's; a dial after it is a double call
        rows7 = open(os.path.join(S, "exports_R05", "E07_2026-11-06.csv"), encoding="utf-8").read().rstrip("\n").split("\n")
        hdr7 = rows7[0].split(","); ix = {k: hdr7.index(k) for k in ("hhid", "ts", "end_ts", "logged_at", "event_id", "prev_hash", "attempt", "appointment_ts", "action")}
        gone7 = mv[1]["hhid"]              # the household that left E07
        def dial_row(ts, eid):
            f = rows7[-1].split(","); f[ix["hhid"]] = gone7; f[ix["ts"]] = ts; f[ix["end_ts"]] = ts[:14] + "01:00"; f[ix["logged_at"]] = ts.replace(" ", "T") + "+08:00"
            f[ix["event_id"]] = eid; f[ix["attempt"]] = "1"; f[ix["appointment_ts"]] = ""; f[ix["action"]] = ""
            return ",".join(f)
        later7 = "\n".join(rows7 + [dial_row("2026-11-06 15:00:00", "00000000-0000-4000-8000-00000000000a"), dial_row("2026-11-13 10:00:00", "00000000-0000-4000-8000-00000000000b")]) + "\n"
        g9 = b.js(f"""(() => {{ const st = LEAD.state(), base = st.imports.map(f => LCORE.readImport(f.name, f.text)).concat([LCORE.readImport('E07_2026-11-13.csv', {json.dumps(later7)})]);
                     const G = LCORE.merge(st.manifest, base, new Map(), RTM.now());
                     const iss = G.issues.filter(i => i.hhid === {json.dumps(gone7)}).map(i => [i.type, i.enumerator, i.detail]);
                     return {{ iss, before: G.dials.some(d => d.event_id === '00000000-0000-4000-8000-00000000000a'), after: G.dials.some(d => d.event_id === '00000000-0000-4000-8000-00000000000b'),
                               holder: LCORE.cases(st.manifest, G).get({json.dumps(gone7)}).enumerator_id }}; }})()""")
        check("after the move, E07's dial on the household before the move is accepted without a question; the one after it is flagged as a double call, kept until the leader decides",
              [t for t, _, _ in g9["iss"]] == ["after_move"] and g9["iss"][0][1] == "E07" and "moved to E02" in g9["iss"][0][2] and g9["before"] and g9["after"] and g9["holder"] == "E02", g9)
        # the leader's session keeps the revised manifest
        before = snapshot()
        b.click('[data-lact="save"]'); b.wait("APP.modal"); b.click('[data-act="modalok"]')
        fpm = wait_new("leader_L2MN-DEMO_R05*.json", before)
        check("the saved leader session holds the revised manifest and logs the moves", json.load(open(fpm, encoding="utf-8"))["manifest"]["revision"] == 2 and any(e["what"] == "households_moved_mid_round" and e["moves"] == 3 for e in json.load(open(fpm, encoding="utf-8"))["events"]))
        # the enumerators open the revised files as updates
        b.click('[data-act="view"][data-v="enum"]')
        b.set_files("#f-assign", [os.path.join(ROOT, "samples", "round05", "E02_assignment_R05.json")]); b.wait("!!APP.S")
        keep02 = next(h["hhid"] for h in json.load(open(os.path.join(ROOT, "samples", "round05", "E02_assignment_R05.json"), encoding="utf-8"))["households"] if h["hhid"] not in moved_ids)
        plan_before = b.js(f"JSON.stringify(APP.C.map.get('{keep02}').h.plan)")
        b.set_files("#f-update", [os.path.join(DL, "E02_assignment_R05.json")]); b.wait("APP.S.events.some(e => e.what === 'assignment_updated')")
        ev = b.js("APP.S.events.find(e => e.what === 'assignment_updated')")
        check("E02 loads the revised file as an update: 3 households added, revision 2 logged, the log kept",
              ev["households_added"] == 3 and ev["households_removed"] == 0 and ev["revision"] == 2 and ev["fingerprint_ok"] == 1, ev)
        check("a household E02 already had keeps its planned first dial and history (carried beside the revised file, which omits them)",
              b.js(f"JSON.stringify(APP.C.map.get('{keep02}').h.plan)") == plan_before and b.js(f"!!APP.S.carry['{keep02}']"))
        prior = b.js(f"(() => {{ const c = APP.C.map.get({json.dumps(mv[0]['hhid'])}); return [c.prior.from, c.nAll, c.state, c.next.kind, RTM.partsIn(c.next.at, APP.C.cal.tz).key]; }})()")
        check("a moved household shows who had it, counts their dials toward the cap and is due today", prior == ["E01", 0, "not_started", "first", "2026-11-13"], prior)
        nev = b.js("APP.S.events.length")
        b.set_files("#f-update", [os.path.join(ROOT, "samples", "round05", "E02_assignment_R05.json")])
        b.wait("APP.toasts.some(t => /older than the assignment in your session/.test(t.text))")
        check("the earlier file, opened again by mistake, is refused as older", b.js("APP.S.events.length") == nev)
        before = snapshot()
        b.click('[data-act="save"]'); wait_new("E02_project_R05*.json", before)
        newest = wait_new("E02_project_R05*.json", before)
        b.click('[data-act="close"]'); b.wait("APP.modal"); b.js("(document.querySelector('[data-act=\"modalalt\"]') || document.querySelector('[data-act=\"modalok\"]')).click(); true"); b.wait("APP.S === null")
        b.set_files("#f-project", [newest]); b.wait("!!APP.S || APP.loadErrors.length")
        check("a session saved after the update reopens with the carried plans", b.js(f"!!APP.S && APP.loadErrors.length === 0 && JSON.stringify(APP.C.map.get('{keep02}').h.plan)") == plan_before, b.js("APP.loadErrors"))
        b.click('[data-act="close"]'); b.wait("APP.modal"); b.js("(document.querySelector('[data-act=\"modalalt\"]') || document.querySelector('[data-act=\"modalok\"]')).click(); true"); b.wait("APP.S === null")
        b.set_files("#f-assign", [os.path.join(ROOT, "samples", "round05", "E07_assignment_R05.json")]); b.wait("!!APP.S")
        b.set_files("#f-update", [os.path.join(DL, "E07_assignment_R05.json")]); b.wait("APP.S.events.some(e => e.what === 'assignment_updated')")
        check("E07 loads the revised file: the moved household is no longer theirs, shown read-only",
              b.js(f"[APP.S.events.find(e => e.what === 'assignment_updated').households_removed, APP.C.map.get({json.dumps(gone7)}).kind]") == [1, "removed"])
        b.click('[data-act="close"]'); b.wait("APP.modal"); b.js("(document.querySelector('[data-act=\"modalalt\"]') || document.querySelector('[data-act=\"modalok\"]')).click(); true"); b.wait("APP.S === null")
        b.click('[data-act="view"][data-v="leader"]')

        # ================= E. exports of KoBoToolbox, Survey Solutions and a calling service (proposal K) =================
        # the same round-5 calls as the enumerators' exports, in the three platforms' layouts (samples/kobo, samples/suso, samples/calls)
        KOBO = os.path.join(ROOT, "samples", "kobo", "rtm_call_record_2026-11-06.csv")
        SUSO_ZIP = os.path.join(ROOT, "samples", "suso", "rtm_call_record_export.zip")
        SVC = os.path.join(ROOT, "samples", "calls", "service_calls_2026-11-06.csv")
        exp_dials = {}
        for r in csv.DictReader(io.StringIO(exp_calls)):
            exp_dials[r["operator"]] = exp_dials.get(r["operator"], 0) + 1
        exp_standalone = [r for r in csv.DictReader(io.StringIO(exp_acts)) if r["via"] != "dial"]

        def merged_now():
            return b.js("""(() => { const g = LEAD.merged(), G = g.G; const by = {};
                G.dials.forEach(d => { by[d.source] = (by[d.source] || 0) + 1; });
                return { dials: by, completes: Object.fromEntries(g.P.list.map(r => [r.enumerator_id, r.complete])), standalone: G.actions.filter(a => a.via !== 'dial').map(a => [a.hhid, a.action, a.target_hhid || '']),
                  files: G.files.map(f => [f.name, f.kind, f.enumerator_id, f.count, f.use]), issues: G.issues.map(i => [i.type, i.status, i.enumerator || '', i.hhid || '', i.severity]),
                  svc: G.serviceStats || null, unverified: [...G.perEnum.values()].every(x => x.unverified), nodur: G.dials.filter(d => d.dur_min === null).length }; })()""")

        start()
        b.js("RTM.setNow('2026-11-07T09:00:00+08:00'); true")
        b.set_files("#lf-hh", [os.path.join(S, "household_list.csv")]); b.wait("!!LEAD.state().inputs.households")
        b.click('[data-lact="tab"][data-t="import"]')
        b.set_files("#lf-manifest", [MANIFEST]); b.wait("!!LEAD.state().manifest")
        txt = b.js("document.querySelector('.panel').innerText")
        check("the Import tab offers a second way in: KoBoToolbox, Survey Solutions or a calling service's exports", "Or from KoBoToolbox, a call-center system, Survey Solutions or a calling service" in txt and b.js("!!document.querySelector('#lf-platform')"))
        b.set_files("#lf-platform", [KOBO]); b.wait("LEAD.state().imports.length === 1")
        k = merged_now()
        check("a KoBoToolbox export (one submission per call, semicolon CSV) is recognised, split by enumerator and merged as unverified CSV exports",
              any(f[1].startswith("RTM call record") and "split into 3 exports" in f[4] for f in k["files"]) and k["unverified"] and sorted(f[2] for f in k["files"] if f[1] == "platform records") == ["E01", "E02", "E07"], k["files"])
        check("the KoBo export gives the same dials per enumerator as the enumerators' own exports", k["dials"] == exp_dials, [k["dials"], exp_dials])
        check("and the same completes per enumerator (Box 3.0: the first complete of 3 minutes or more)", k["completes"] == py_cmp, [k["completes"], py_cmp])
        check("a record without a call (result NONE) becomes an action without a dial, as in the enumerators' own logs",
              sorted(k["standalone"]) == sorted([a["hhid"], a["action"], a["target_hhid"]] for a in exp_standalone), [k["standalone"], [(a["hhid"], a["action"]) for a in exp_standalone]])
        check("no merge question is open: the form's dials sit on their own households (each part is noted as unverified, as a CSV without its JSON is)",
              not any(i[1] == "open" and i[4] == "action" for i in k["issues"]) and sum(1 for i in k["issues"] if i[0] == "unverified") == 3, [i for i in k["issues"] if i[1] == "open"])
        kobo_rows = list(csv.DictReader(open(KOBO, encoding="utf-8"), delimiter=";"))
        blank_ends = sum(1 for r in kobo_rows if r["result"] not in ("NONE",) and r["call_end"] == "")
        check("call ends the enumerator left blank give dials without a length", k["nodur"] == blank_ends and blank_ends > 0, [k["nodur"], blank_ends])
        willing = [r for r in kobo_rows if r["willing"] == "later"]
        check("the willingness to take part travels into the merged log's note",
              willing and b.js(f"LEAD.merged().G.dials.filter(d => (d.note || '').includes('willing: later')).length") == len([r for r in willing if r["result"] != "NONE"]), len(willing))
        # the calling service's log fills the blanks and reports what nobody recorded
        b.set_files("#lf-platform", [SVC]); b.wait("LEAD.state().imports.length === 2")
        s = merged_now()
        svc_rows = list(csv.DictReader(open(SVC, encoding="utf-8")))
        check("the calling service's log (Twilio's call fields) is matched to the recorded calls by number and time; calls to numbers outside the list are counted apart",
              s["svc"] and s["svc"]["calls"] == len(svc_rows) and s["svc"]["unknown"] == 2 and s["svc"]["matched"] + s["svc"]["unlogged"] + s["svc"]["unknown"] == len(svc_rows), s["svc"])
        check("service calls that nobody recorded are listed for the leader; the matched ones fill the missing call ends",
              3 <= s["svc"]["unlogged"] <= 6 and sum(1 for i in s["issues"] if i[0] == "unlogged_call") == s["svc"]["unlogged"] and s["svc"]["filled"] >= 1 and s["nodur"] == blank_ends - s["svc"]["filled"], s["svc"])
        check("filled call ends change no complete (completes keep their own length)", s["completes"] == py_cmp)
        b.shot(os.path.join(SCREENS, "leader_11_platform.png"), full=True)
        # the Survey Solutions export: the zip is opened in the browser, the roster joined with its main file
        b.click('[data-lact="clearimports"]'); b.wait("LEAD.state().imports.length === 0")
        b.set_files("#lf-platform", [SUSO_ZIP]); b.wait("LEAD.state().imports.length === 2")
        u = merged_now()
        check("a Survey Solutions export zip is opened in the browser: its main file and the calls roster are joined by interview__id and split by enumerator",
              any("joined with" in f[4] and "split into 3 exports" in f[4] for f in u["files"]) and sorted(f[2] for f in u["files"] if f[1] == "Survey Solutions records") == ["E01", "E02", "E07"], u["files"])
        check("a roster without its main file is reported, not merged",
              b.js(f"""(() => {{ const st = LEAD.state(), calls = st.imports.find(f => f.name.endsWith('calls.tab'));
                  const G = LCORE.merge(st.manifest, [LCORE.readImport(calls.name, calls.text)], new Map(), RTM.now());
                  return G.files.length === 1 && G.files[0].use.startsWith('not merged') && G.dials.length === 0; }})()"""))
        # ---- the two real layouts (proposal K.2): the L2PHL form's export, the L2CU call-center log, sample table and interview form ----
        PHL = os.path.join(ROOT, "samples", "kobo", "L2MN-DEMO_CATI_R05_-_all_versions_-_False_-_2026-11-06-18-00-00.csv")
        CATI = os.path.join(ROOT, "samples", "cati", "calllog_2026-11-06.csv")
        SAMPLE = os.path.join(ROOT, "samples", "cati", "sample_2026-11-06.csv")
        IV = os.path.join(ROOT, "samples", "kobo", "L2MN-DEMO_CATI_R05_interviews_-_latest_version_-_False_-_2026-11-06-18-00-00.csv")
        b.click('[data-lact="clearimports"]'); b.wait("LEAD.state().imports.length === 0")
        b.set_files("#lf-platform", [PHL]); b.wait("LEAD.state().imports.length === 1")
        q = merged_now()
        check("a KoBoToolbox export in the L2PHL layout (the interview form that records each attempt: int_id, hhid, Z8, X15) is recognised by its columns and read under the tool's IDs through the manifest's platform codes",
              any(f[1].startswith("L2PHL") for f in q["files"]) and sorted(f[2] for f in q["files"] if f[1] == "platform records") == ["E01", "E02", "E07"], q["files"])
        check("it gives the same dials and completes per enumerator as the enumerators' own exports (X15 code 1 is a complete)", q["dials"] == exp_dials and q["completes"] == py_cmp, [q["dials"], exp_dials, q["completes"], py_cmp])
        check("no merge question is open for it", not any(i[1] == "open" and i[4] == "action" for i in q["issues"]), [i for i in q["issues"] if i[1] == "open"])
        b.click('[data-lact="clearimports"]'); b.wait("LEAD.state().imports.length === 0")
        b.set_files("#lf-platform", [CATI]); b.wait("LEAD.state().imports.length === 1")
        c = merged_now()
        cati_rows = list(csv.DictReader(open(CATI, encoding="utf-8"), delimiter=";"))
        other = sum(1 for r in cati_rows if r["round"] != "5")
        check("a call-center system's log in the L2CU layout (semicolon CSV, cumulative across rounds) is recognised; the rows of other rounds are left out",
              any(f[1].startswith("Call-center system log (L2CU") and f"{other} rows of other rounds left out" in f[4] for f in c["files"]), c["files"])
        check("its dials and completes per enumerator match the enumerators' own exports (CS is a complete; the operator code CATI-07 is read as E07 through the platform code CATI07; the round's time zone supplies the missing offsets)",
              c["dials"] == exp_dials and c["completes"] == py_cmp, [c["dials"], c["completes"]])
        slots = b.js("(() => { const G = LEAD.merged().G; return [G.dials.length, G.dials.filter(d => d.phone_slot).length, G.dials.filter(d => /\+08:00$/.test(d.ts)).length]; })()")
        check("the number dialed becomes the household's slot, from the household list, and is not kept", slots[0] == slots[1] and slots[2] == slots[0], slots)
        b.set_files("#lf-platform", [IV]); b.wait("LEAD.state().imports.length === 2")
        v = merged_now()
        iv_stats = b.js("LEAD.merged().G.ivStats")
        check("the interview form's export (L2CU layout: hhid, ecode, agreement) is checked against the logged calls: two interviews without a completed call and two completed calls without an interview are listed",
              bool(iv_stats) and iv_stats["unlogged"] == 2 and iv_stats["without"] == 2 and iv_stats["matched"] == sum(py_cmp.values()) - 2
              and sum(1 for i in v["issues"] if i[0] == "interview_unlogged") == 2 and sum(1 for i in v["issues"] if i[0] == "complete_without_interview") == 2, iv_stats)
        key = b.js("LEAD.merged().G.issues.find(i => i.type === 'interview_unlogged').key")
        b.click(f'[data-lact="resolve"][data-key="{key}"][data-v="count"]'); b.wait("LEAD.state().resolutions.length >= 1")
        v2 = merged_now()
        check("an interview the leader counts becomes a complete", sum(v2["completes"].values()) == sum(v["completes"].values()) + 1, [v["completes"], v2["completes"]])
        # the system's sample table: its view of the assignment against the manifest
        smp_rows = list(csv.DictReader(open(SAMPLE, encoding="utf-8"), delimiter=";"))
        sys_rep = next(r for r in smp_rows if r["replacement_for"])
        b.set_files("#lf-platform", [SAMPLE]); b.wait("LEAD.state().imports.length === 3")
        w = merged_now()
        smp_stats = b.js("LEAD.merged().G.sampleStats")
        check("the system's sample table (L2CU layout) is read against the manifest: the one household the system has with another operator is listed for a decision, and the replacement the system activated is counted as activated",
              bool(smp_stats) and smp_stats["differing"] == 1 and smp_stats["activated"] == 1 and sum(1 for i in w["issues"] if i[0] == "system_assignment" and i[1] == "open") == 1
              and sum(1 for i in w["issues"] if i[0] == "system_replacement") == 1 and b.js(f"LEAD.merged().C.has({json.dumps(sys_rep['hhid'])})"), [smp_stats, sys_rep["hhid"]])
        b.click('[data-lact="clearimports"]'); b.wait("LEAD.state().imports.length === 0")
        b.set_files("#lf-platform", [IV]); b.wait("LEAD.state().imports.length === 1")
        v3 = merged_now()
        check("without a call log, the interview export is read as the completes and refusals", any("read as completes" in f[4] for f in v3["files"]) and sum(v3["completes"].values()) == sum(py_cmp.values()), [v3["files"], v3["completes"]])
        # a file with unknown columns is mapped once, in the Import tab
        unk = os.path.join(SCRATCH, "unknown_log.csv")
        with open(unk, "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f); w.writerow(["op", "household", "t0", "t1", "res", "rid"])
            for r in cati_rows:
                if r["round"] == "5":
                    w.writerow([r["operator"], r["hhid"], r["start_time"], r["end_time"], r["status"], r["callid"]])
        b.click('[data-lact="clearimports"]'); b.wait("LEAD.state().imports.length === 0")
        b.set_files("#lf-platform", [unk]); b.wait("LEAD.state().imports.length === 1")
        check("a table with unknown columns opens the mapping card instead of being merged", b.js("!!document.querySelector('#lmap')") and merged_now()["dials"] == {}, merged_now()["files"])
        b.js("""(() => { const sel = (k, v) => { const s = document.querySelector(`[data-lprof="${k}"]`); if (!s) return; s.value = v; s.dispatchEvent(new Event('change', { bubbles: true })); };
            sel('col:hhid', 'household'); sel('col:enumerator', 'op'); sel('col:start', 't0'); sel('col:end', 't1'); sel('col:result', 'res'); sel('col:id', 'rid');
            const codes = { CS: 'COMPLETE', NA: 'NO_ANSWER', NT: 'NOT_AVAILABLE', NWN: 'DEAD_NUMBER', WN: 'WRONG_NUMBER', OSA: 'SWITCHED_OFF', IS: 'PARTIAL', BS: 'BUSY', HU: 'OTHER', NC: 'NONE' };
            Object.entries(codes).forEach(([k, v]) => sel(`code:${k}`, v)); sel('name', 'Our call-center log'); return true; })()""")
        b.click('[data-lact="saveprofile"]'); b.wait("LEAD.state().profiles.length === 1")
        u2 = merged_now()
        check("once the leader has mapped its columns and codes, the file merges like a known layout, and the mapping is kept in the session",
              u2["dials"] == exp_dials and u2["completes"] == py_cmp and b.js("LEAD.state().profiles[0].name") == "Our call-center log", [u2["dials"], u2["completes"], u2["files"]])
        b.shot(os.path.join(SCREENS, "leader_12_mapping.png"), full=True)
        check("the Survey Solutions export gives the same dials and completes per enumerator as the KoBo one", u["dials"] == exp_dials and u["completes"] == py_cmp, [u["dials"], u["completes"]])
        # the platform files the leader hands to the platforms, from the manifest
        b.click('[data-lact="tab"][data-t="calllog"]')
        before = snapshot()
        b.click('[data-lact="export"][data-what="kobo"]'); b.click('[data-lact="export"][data-what="suso"]')
        b.click('[data-lact="export"][data-what="choices"]'); b.click('[data-lact="export"][data-what="cati"]')
        fk = wait_new("households*.csv", before); fs = wait_new("assignments*.tab", before); fc = wait_new("choices_households*.csv", before); fa = wait_new("cati_assignments*.csv", before)
        hk = list(csv.DictReader(open(fk, encoding="utf-8")))
        sample_hk = {r["hhid"]: r for r in csv.DictReader(open(os.path.join(ROOT, "samples", "kobo", "households.csv"), encoding="utf-8"))}
        same = all(r["hhid"] in sample_hk and all(r[c] == sample_hk[r["hhid"]][c] for c in ("enumerator_id", "phone_1", "phone_2", "phone_3", "kind", "replaces", "rank", "program", "round")) for r in hk)
        check("households.csv for the KoBo form: one row per household and reserve, with the enumerator, numbers and replacement links of the manifest (the sample file, apart from the plan)",
              len(hk) == len(sample_hk) and same and "name" in hk[0] and "label" in hk[0], [len(hk), len(sample_hk)])
        ta = [l.split("\t") for l in open(fs, encoding="utf-8").read().rstrip("\n").split("\n")]
        check("assignments.tab for Survey Solutions: the manifest's households with _responsible and _quantity",
              ta[0][-2:] == ["_responsible", "_quantity"] and len(ta) - 1 == len(man["households"]) and all(r[-1] == "1" and r[-2] in ("E01", "E02", "E07") for r in ta[1:]), [ta[0], len(ta)])
        ch = list(csv.DictReader(open(fc, encoding="utf-8")))
        check("choices_households.csv (the L2PHL form's SN list): one row per household and reserve, the enumerator's platform code in filter_int_id",
              len(ch) == len(sample_hk) and all(r["list_name"] == "SN" and r["filter_int_id"] == "CATI" + sample_hk[r["name"]]["enumerator_id"][1:] for r in ch), [len(ch), ch[:1]])
        ca = list(csv.DictReader(open(fa, encoding="utf-8"), delimiter=";"))
        check("cati_assignments.csv (the call-center system's table, in the L2CU sample table's names): household, catiid, planned first dial and numbers for every household and reserve",
              len(ca) == len(sample_hk) and all(r["catiid"].startswith("CATI") and r["hhid"] in sample_hk and r["respondent_phone"] == sample_hk[r["hhid"]]["phone_1"] for r in ca) and "date_nextcall" in ca[0], [len(ca), ca[:1]])

        # phone width
        b.ws.call("Emulation.setDeviceMetricsOverride", width=390, height=844, deviceScaleFactor=2, mobile=True)
        b.click('[data-lact="tab"][data-t="progress"]')
        b.shot(os.path.join(SCREENS, "leader_08_phone.png"))
        check("at phone width the page has no sideways scroll (tables scroll inside their cards)", b.js("document.documentElement.scrollWidth <= window.innerWidth + 1"), b.js("[document.documentElement.scrollWidth, window.innerWidth]"))
        b.ws.call("Emulation.clearDeviceMetricsOverride")
    finally:
        b.close()
    passed = sum(ok for _, ok, _ in results)
    lines = [f"Leader view tests, {time.strftime('%d %b %Y %H:%M')}: {passed} of {len(results)} passed"] + [("PASS " if ok else "FAIL ") + n for n, ok, _ in results]
    open(os.path.join(HERE, "results_leader.txt"), "w").write("\n".join(lines) + "\n")
    print(lines[0])
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
